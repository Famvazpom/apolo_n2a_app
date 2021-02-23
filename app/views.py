# -*- encoding: utf-8 -*-
"""
Copyright (c) 2019 - present AppSeed.us
"""

from django.contrib.auth.decorators import login_required
from django.shortcuts import render, get_object_or_404, redirect
from django.template import loader
from django.http import HttpResponse,StreamingHttpResponse
from django.views.generic.base import TemplateView
from django import template

from werkzeug.wsgi import FileWrapper
import pandas as pd
import os
import json
import io
import csv
import docx

class resultsView(TemplateView):
    template_name = 'results.html'

    def get_file_data(self,file):
        with open(file) as json_file:
            return json.load(json_file)
    
    def process_dataframe(self,file):
        array = []
        df = pd.read_json(file)
        for i,values in df.iterrows():
            row = []
            for j in values:
                row.append(str(j).replace('./bd/',''))
            array.append(row)
        return array

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['directorios'] = self.get_file_data('app/db_results/directorios.json')
        context["palabras"] = self.process_dataframe('app/db_results/palabras_ciclo.json')

        context['ocurrencias'] = self.process_dataframe('app/db_results/ocurrencia.json')
        context['unicas'] = self.process_dataframe('app/db_results/unicas.json')
        context['values'] = self.process_dataframe('app/db_results/values.json')
        return context
    
    def get(self,request,*args, **kwargs):
        return render(request,self.template_name,self.get_context_data())

class export(TemplateView):

    def read_data(self):
        return pd.read_json('app/db_results/values.json')


    def csv(self):
        data = self.read_data()
        print(data)
        buffer = io.BytesIO() 
        
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = "attachment; filename=Materias.csv"
        data.to_csv(response)
        return response
    
    def xls(self):
        data = self.read_data()
        buffer = io.BytesIO()
        response =  HttpResponse(content_type='application/vnd.ms-excel')
        response['Content-Disposition'] = "attachment; filename=Materias.xls"
        data.to_excel(response)
        return response


    def doc(self):
        df = self.read_data()
        buffer = io.BytesIO()
        doc = docx.Document()

        # add a table to the end and create a reference variable
        # extra row is so we can add the header row
        t = doc.add_table(df.shape[0]+1, df.shape[1])

        # add the header rows.
        for j in range(df.shape[-1]):
            t.cell(0,j).text = df.columns[j]

        # add the rest of the data frame
        for i in range(df.shape[0]):
            for j in range(df.shape[-1]):
                t.cell(i+1,j).text = str(df.values[i,j])

        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        response['Content-Disposition'] = "attachment; filename=Materias.doc"
        doc.save(response)
        
        return response

    def html(self):
        data = self.read_data()
        buffer = io.BytesIO()
        response = HttpResponse(content_type='text/plain')
        response['Content-Disposition'] = "attachment; filename=Materias.html"
        data.to_html(response)
        return response

    def txt(self):
        data = self.read_data()
        buffer = io.BytesIO()
        response = HttpResponse(content_type='text/plain')
        response['Content-Disposition'] = "attachment; filename=Materias.txt"
        data.to_csv(response)
        return response

    def get(self,request,etype,*args, **kwargs):
        if etype == 0:
            return self.csv()
        elif etype == 1:
            return self.xls()
        elif etype == 2:
            return self.doc()
        elif etype == 3:
            return self.html()
        elif etype == 4:
            return self.txt()